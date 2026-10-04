"""Revision 2 of the RehletShifaa Users, Roles and Access Guide.

Reuses the drawing and document helpers of build_manual.py but writes to new file names and a new
asset folder, so the original manual and its figures are never overwritten.
"""
from pathlib import Path

from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx import Document
from PIL import Image, ImageDraw

import build_manual as bm
from build_manual import (
    NAVY, TEAL, BLUE, GOLD, GREEN, RED, INK, MID, LIGHT, PALE, WHITE,
    F_TITLE, F_H1, F_H2, F_BODY, F_SMALL, F_BOLD, F_CARD,
    hexrgb, rounded, arrow, canvas,
    set_table_borders, set_cell_shading, set_cell_margins, set_cell_width,
    font_run, style_paragraph, add_body, add_bullets, add_numbered, add_heading,
    add_caption, add_picture, add_page_break, add_page_number,
    set_repeat_table_header, set_row_cant_split,
)

OUT = bm.OUT
ASSETS = OUT / "users-roles-manual-rev2-assets"
DOCX = OUT / "rehletshifaa-users-and-roles-manual-rev2.docx"
bm.ASSETS = ASSETS  # bm.save() writes into the revision folder

AMBER_FILL = "FFF6E0"
RED_FILL = "F8E9EA"
GREEN_FILL = "EDF5EF"


def draw_wrapped(d, xy, text, fnt, fill, width_chars, spacing=5, anchor=None):
    """Centre-aligned variant of build_manual.draw_wrapped."""
    from textwrap import wrap
    lines = []
    for paragraph in text.split("\n"):
        lines.extend(wrap(paragraph, width=width_chars) or [""])
    align = "center" if anchor and anchor[0] == "m" else "left"
    d.multiline_text(xy, "\n".join(lines), font=fnt, fill=fill, spacing=spacing, anchor=anchor, align=align)


def add_table(doc, headers, rows, widths=None, font_size=9.2, header_fill=NAVY):
    """build_manual.add_table with real grid widths (LibreOffice ignores cell widths alone) and left-aligned text."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if widths:
        for i, w in enumerate(widths):
            table.columns[i].width = Inches(w)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, 110, 115, 110, 115)
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        style_paragraph(p, after=0, line=1.0)
        p.paragraph_format.keep_with_next = True
        font_run(p.add_run(text), font_size, True, WHITE)
    for ri, row in enumerate(rows):
        new_row = table.add_row()
        set_row_cant_split(new_row)
        for i, text in enumerate(row):
            cell = new_row.cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, 100, 115, 100, 115)
            if widths:
                set_cell_width(cell, widths[i])
            if ri % 2 == 1:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            style_paragraph(p, after=0, line=1.02)
            font_run(p.add_run(str(text)), font_size, i == 0 and len(headers) > 2 and len(str(text)) <= 6 and any(c.isdigit() for c in str(text)), INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def save(img, name):
    path = ASSETS / name
    img.save(path, quality=95)
    return path


# --------------------------------------------------------------------------- diagrams

def diagram_populations():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 52), "Who uses RehletShifaa", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    items = [
        ("Platform governance", "Platform Account Owner\nSystem Administrators", NAVY),
        ("RehletShifaa teams", "Care coordination\nTravel and logistics\nCommercial and finance\nProvider operations\nCredentialing and review", TEAL),
        ("Provider organizations", "Organization Owners\nPractice Managers\nConsultants\nAssociate Doctors\nConsultant Assistants", BLUE),
        ("Patients", "Patients\nPatient Representatives", GREEN),
    ]
    for (title, body, color), x in zip(items, [60, 450, 840, 1230]):
        rounded(d, (x, 180, x + 310, 720), LIGHT if color != GREEN else GREEN_FILL, color, 4)
        d.ellipse((x + 105, 215, x + 205, 315), fill=hexrgb(color))
        d.text((x + 155, 265), title[0], font=F_TITLE, fill="white", anchor="mm")
        d.text((x + 155, 350), title, font=F_H2, fill=hexrgb(INK), anchor="ma")
        draw_wrapped(d, (x + 155, 440), body, F_BODY, hexrgb(INK), 24, 13, "ma")
    d.text((800, 820), "One platform, four user populations, clear boundaries", font=F_H2, fill=hexrgb(MID), anchor="mm")
    return save(img, "figure-01-user-populations.png")


def diagram_governance():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 45), "Platform governance and day-to-day authority", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    rounded(d, (500, 120, 1100, 245), "F3E8C9", GOLD, 4)
    d.text((800, 160), "Platform Account Owner", font=F_H1, fill=hexrgb(INK), anchor="ma")
    d.text((800, 212), "Approves administrator changes and transfers ownership", font=F_SMALL, fill=hexrgb(INK), anchor="ma")
    arrow(d, (800, 250), (800, 320), NAVY)
    rounded(d, (500, 325, 1100, 450), LIGHT, NAVY, 4)
    d.text((800, 362), "System Administrators", font=F_H1, fill=hexrgb(INK), anchor="ma")
    d.text((800, 414), "Invite staff, change jobs and grant access", font=F_SMALL, fill=hexrgb(INK), anchor="ma")
    for x in [230, 545, 860, 1175]:
        arrow(d, (800, 455), (x, 555), TEAL, 4)
    labels = [
        "Care delivery jobs\nCoordination, travel, finance",
        "Provider governance\nProvider operations, credentialing",
        "Content and identity\nCare journeys, patient identity",
        "Independent oversight\nCompliance audit",
    ]
    for x, label in zip([80, 395, 710, 1025], labels):
        rounded(d, (x, 560, x + 300, 760), PALE, TEAL, 3)
        draw_wrapped(d, (x + 150, 660), label, F_BODY, hexrgb(INK), 24, 10, "mm")
    d.text((800, 830), "Neither position gives access to patient cases or medical documents", font=F_BOLD, fill=hexrgb(RED), anchor="mm")
    return save(img, "figure-02-platform-governance.png")


def diagram_ownership_transfer():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 45), "Transferring platform ownership", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    steps = [
        (70, "1", "Owner nominates a successor", "Strong sign-in, recent sign-in and a written reason"),
        (450, "2", "Successor is set up", "Strong sign-in becomes mandatory for the successor"),
        (830, "3", "Successor accepts", "Signs in with strong sign-in and accepts the nomination"),
        (1210, "4", "Ownership moves", "Previous ownership ends. History is kept"),
    ]
    for x, num, title, body in steps:
        d.ellipse((x + 110, 140, x + 210, 240), fill=hexrgb(GOLD))
        d.text((x + 160, 190), num, font=F_TITLE, fill="white", anchor="mm")
        rounded(d, (x, 280, x + 320, 560), PALE, GOLD, 3)
        draw_wrapped(d, (x + 160, 310), title, F_CARD, hexrgb(INK), 22, 6, "ma")
        draw_wrapped(d, (x + 160, 400), body, F_BODY, hexrgb(INK), 22, 8, "ma")
    for x in [290, 670, 1050]:
        arrow(d, (x, 190), (x + 150, 190), GOLD, 5)
    rounded(d, (70, 610, 1530, 850), RED_FILL, RED, 3)
    lines = [
        "A System Administrator can never transfer ownership.",
        "Only one transfer can be pending. A repeated submission does not create a second one.",
        "Owner unreachable and no administrator left: vendor recovery with a 72-hour waiting period that can be cancelled.",
    ]
    for i, line in enumerate(lines):
        d.text((110, 640 + i * 70), line, font=F_BODY, fill=hexrgb(INK))
    return save(img, "figure-03-ownership-transfer.png")


def diagram_provider_tree():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 40), "One person across organizations, clinics and roles", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    rounded(d, (545, 100, 1055, 200), GREEN_FILL, GREEN, 4)
    d.text((800, 132), "Dr Sara Haddad", font=F_H1, fill=hexrgb(INK), anchor="ma")
    d.text((800, 178), "One sign-in and one practitioner profile", font=F_SMALL, fill=hexrgb(INK), anchor="ma")
    orgs = [
        (80, "Al Noor Healthcare Group", "Organization Owner (all clinics)\nConsultant at Dubai and Sharjah"),
        (575, "Gulf Heart Center", "Practice Manager at Gulf Heart Main only"),
        (1070, "Crescent Medical", "Associate Doctor at Crescent Rehab\nsupervised by Dr Lina Mansour"),
    ]
    for x, title, body in orgs:
        arrow(d, (800, 205), (x + 225, 315), BLUE, 4)
        rounded(d, (x, 320, x + 450, 520), LIGHT, BLUE, 4)
        d.text((x + 225, 350), title, font=F_H2, fill=hexrgb(INK), anchor="ma")
        draw_wrapped(d, (x + 225, 410), body, F_BODY, hexrgb(INK), 34, 8, "ma")
    for x, label in [(80, "Dubai Clinic\nSharjah Clinic"), (575, "Gulf Heart Main"), (1070, "Crescent Rehab")]:
        arrow(d, (x + 225, 525), (x + 225, 610), TEAL, 4)
        rounded(d, (x + 55, 615, x + 395, 760), PALE, TEAL, 3)
        draw_wrapped(d, (x + 225, 688), label, F_H2, hexrgb(INK), 24, 8, "mm")
    d.text((800, 835), "Each membership, clinic link and role starts and ends on its own", font=F_BOLD, fill=hexrgb(MID), anchor="mm")
    return save(img, "figure-04-multi-organization.png")


def diagram_multi_clinic():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 30), "One organization, several clinics: who reaches where", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    rounded(d, (560, 95, 1540, 175), "F3E8C9", GOLD, 4)
    d.text((1050, 135), "Al Noor Healthcare Group (one organization, one contract)", font=F_H2, fill=hexrgb(INK), anchor="mm")
    cols = [580, 820, 1060, 1300]
    facs = ["Dubai Clinic\n(default clinic)", "Sharjah Clinic", "Abu Dhabi\nHospital", "Al Noor Online\n(virtual care)"]
    for x, label in zip(cols, facs):
        arrow(d, (x + 110, 178), (x + 110, 215), GOLD, 3)
        rounded(d, (x, 220, x + 220, 320), PALE, TEAL, 3)
        draw_wrapped(d, (x + 110, 270), label, F_SMALL, hexrgb(INK), 18, 4, "mm")
    rows = [
        ("Organization Owner", "Organization-wide", 0, 4, NAVY),
        ("Practice Manager", "Dubai and Sharjah only", 0, 2, BLUE),
        ("Consultant (Dr Sara)", "Linked to Dubai and Sharjah", 0, 2, TEAL),
        ("Virtual Consultant", "Al Noor Online only", 3, 4, GREEN),
        ("RehletShifaa credentialing", "Coverage of Dubai only", 0, 1, GOLD),
    ]
    y = 360
    for name, note, start, end, color in rows:
        d.text((60, y + 8), name, font=F_BOLD, fill=hexrgb(INK))
        d.text((60, y + 42), note, font=F_SMALL, fill=hexrgb(MID))
        for x in cols:
            d.rounded_rectangle((x, y, x + 220, y + 70), radius=12, outline=hexrgb("D9D9D9"), width=2)
        rounded(d, (cols[start] + 8, y + 12, cols[end - 1] + 212, y + 58), color, color, 2, 16)
        d.text(((cols[start] + cols[end - 1] + 220) // 2, y + 35), "access", font=F_SMALL, fill="white", anchor="mm")
        y += 90
    d.text((800, 850), "The default clinic pre-fills forms. It never grants access by itself.", font=F_BOLD, fill=hexrgb(RED), anchor="mm")
    return save(img, "figure-05-multi-clinic.png")


def diagram_two_paths():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 35), "Two separate ways to reach a provider organization", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    lanes = [
        (60, BLUE, "Provider users", [
            "Provider membership\nin one organization",
            "Clinic links\n(where the person works)",
            "Role and scope\nwhole organization, selected clinics,\nmanaged clinicians or self",
        ], "Accepted by the person through an offer.\nManaged by the organization."),
        (840, TEAL, "RehletShifaa internal staff", [
            "Staff record\none job plus extra responsibilities",
            "RehletShifaa coverage\nof an organization or selected clinics",
            "No clinic link needed.\nNot a provider employee.",
        ], "Granted and ended only by a\nRehletShifaa System Administrator."),
    ]
    for x, color, title, boxes, foot in lanes:
        rounded(d, (x, 100, x + 700, 790), "FFFFFF", color, 4)
        d.text((x + 350, 125), title, font=F_H1, fill=hexrgb(color), anchor="ma")
        y = 190
        for i, text in enumerate(boxes):
            rounded(d, (x + 60, y, x + 640, y + 130), PALE, color, 3)
            draw_wrapped(d, (x + 350, y + 65), text, F_BODY, hexrgb(INK), 40, 6, "mm")
            if i < len(boxes) - 1:
                arrow(d, (x + 350, y + 132), (x + 350, y + 168), color, 4)
            y += 170
        draw_wrapped(d, (x + 350, 735), foot, F_SMALL, hexrgb(MID), 44, 6, "mm")
    d.text((800, 845), "Providers never see, change or remove RehletShifaa coverage", font=F_BOLD, fill=hexrgb(RED), anchor="mm")
    return save(img, "figure-06-two-paths.png")


def diagram_context_switch():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 40), "Signing in once and switching context", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    rounded(d, (50, 160, 380, 440), GREEN_FILL, GREEN, 4)
    draw_wrapped(d, (215, 300), "Sign in once\nwith strong sign-in\nwhen any of your\nroles needs it", F_BODY, hexrgb(INK), 20, 8, "mm")
    arrow(d, (385, 300), (455, 300), BLUE, 5)
    rounded(d, (460, 120, 820, 480), LIGHT, BLUE, 4)
    d.text((640, 145), "Choose organization", font=F_CARD, fill=hexrgb(INK), anchor="ma")
    for i, name in enumerate(["Al Noor Healthcare Group", "Gulf Heart Center", "Crescent Medical"]):
        fill = "FFFFFF" if i else "D6E8F2"
        rounded(d, (490, 200 + i * 85, 790, 265 + i * 85), fill, BLUE, 2, 12)
        d.text((640, 232 + i * 85), name, font=F_SMALL, fill=hexrgb(INK), anchor="mm")
    arrow(d, (825, 300), (895, 300), BLUE, 5)
    rounded(d, (900, 120, 1180, 480), LIGHT, TEAL, 4)
    d.text((1040, 145), "Choose clinic", font=F_CARD, fill=hexrgb(INK), anchor="ma")
    for i, name in enumerate(["Dubai Clinic", "Sharjah Clinic"]):
        fill = "FFFFFF" if i else "D3ECEC"
        rounded(d, (930, 200 + i * 85, 1150, 265 + i * 85), fill, TEAL, 2, 12)
        d.text((1040, 232 + i * 85), name, font=F_SMALL, fill=hexrgb(INK), anchor="mm")
    arrow(d, (1185, 300), (1255, 300), BLUE, 5)
    rounded(d, (1260, 120, 1560, 480), "F3E8C9", GOLD, 4)
    draw_wrapped(d, (1410, 300), "Server re-checks\nevery action:\nmembership, clinic,\nrole, status and\neligibility", F_BODY, hexrgb(INK), 20, 8, "mm")
    rounded(d, (50, 540, 1560, 820), RED_FILL, RED, 3)
    notes = [
        "The lists show only organizations and clinics where the person has an active relationship.",
        "The selector changes what you see. It never grants access.",
        "Opening another organization's page by link or bookmark is refused if you have no role there.",
    ]
    for i, line in enumerate(notes):
        d.text((90, 570 + i * 78), line, font=F_BODY, fill=hexrgb(INK))
    return save(img, "figure-07-context-switch.png")


def diagram_offer_flow():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 45), "Joining a provider organization needs the person's consent", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    steps = [
        (60, "1", "Offer is sent", "Email is used only to find the person"),
        (370, "2", "Identity is found", "One account is reused or created once"),
        (680, "3", "Person reviews", "Organization, clinics, roles and dates"),
        (990, "4", "Accept or decline", "Manager roles need strong sign-in"),
        (1300, "5", "Access starts", "Membership, clinics and roles start together"),
    ]
    for x, num, title, body in steps:
        d.ellipse((x, 190, x + 120, 310), fill=hexrgb(TEAL))
        d.text((x + 60, 250), num, font=F_TITLE, fill="white", anchor="mm")
        rounded(d, (x - 35, 350, x + 155, 660), PALE, TEAL, 3)
        draw_wrapped(d, (x + 60, 380), title, F_CARD, hexrgb(INK), 14, 5, "ma")
        draw_wrapped(d, (x + 60, 480), body, F_BODY, hexrgb(INK), 14, 7, "ma")
    for x in [180, 490, 800, 1110]:
        arrow(d, (x, 250), (x + 150, 250), BLUE, 5)
    d.text((800, 725), "Nothing is active before acceptance", font=F_H1, fill=hexrgb(RED), anchor="mm")
    d.text((800, 790), "The inviter always sees \"Invitation sent\", whether or not the person already has an account", font=F_BOLD, fill=hexrgb(MID), anchor="mm")
    d.text((800, 845), "Declining or letting one offer expire changes nothing in other organizations", font=F_BOLD, fill=hexrgb(MID), anchor="mm")
    return save(img, "figure-08-provider-offer.png")


def diagram_lifecycle():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 40), "Local actions stay local", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    rounded(d, (560, 105, 1040, 205), GREEN_FILL, GREEN, 4)
    d.text((800, 138), "One sign-in identity", font=F_H1, fill=hexrgb(INK), anchor="ma")
    d.text((800, 182), "Stays enabled after any local action", font=F_SMALL, fill=hexrgb(INK), anchor="ma")
    cards = [
        (40, "Clinic link ends", "That clinic only.\nOther clinics and\norganizations continue.", BLUE),
        (345, "Membership ends", "That organization only.\nOther organizations\ncontinue.", BLUE),
        (650, "Clinic suspended", "Everyone is stopped\nat that clinic only.\nSibling clinics continue.", BLUE),
        (955, "Organization suspended", "All its clinics stop.\nOther organizations\ncontinue.", BLUE),
        (1260, "Sign-in disabled", "Everything stops.\nOnly for compromise or\nan authorized global action.", RED),
    ]
    for x, title, body, color in cards:
        arrow(d, (800, 210), (x + 150, 330), color, 4)
        rounded(d, (x, 335, x + 300, 660), PALE, color, 4)
        draw_wrapped(d, (x + 150, 370), title, F_CARD, hexrgb(INK), 22, 6, "ma")
        draw_wrapped(d, (x + 150, 450), body, F_BODY, hexrgb(INK), 24, 8, "ma")
    d.text((800, 735), "RehletShifaa coverage and provider memberships never affect each other", font=F_BOLD, fill=hexrgb(MID), anchor="mm")
    d.text((800, 810), "Removing a person from one organization never disables their sign-in", font=F_H2, fill=hexrgb(RED), anchor="mm")
    return save(img, "figure-09-lifecycle-isolation.png")


def diagram_pricing():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 45), "How the platform chooses a service price", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    levels = [
        (120, 135, 1480, 245, LIGHT, BLUE, "1. This clinician at this clinic", "Most specific. Used first when it exists"),
        (200, 270, 1400, 380, PALE, TEAL, "2. This clinic", "Used when there is no clinician-at-clinic price"),
        (280, 405, 1320, 515, LIGHT, NAVY, "3. This clinician, organization-wide", "Used when no clinic-level price exists"),
        (360, 540, 1240, 650, PALE, GOLD, "4. Organization default", "Used when nothing more specific exists"),
        (440, 675, 1160, 785, LIGHT, GREEN, "5. Catalogue fallback", "Used only when levels 1 to 4 are absent"),
    ]
    for x1, y1, x2, y2, fill, color, title, body in levels:
        rounded(d, (x1, y1, x2, y2), fill, color, 4)
        d.text((x1 + 35, y1 + 28), title, font=F_H2, fill=hexrgb(INK))
        d.text((x1 + 35, y1 + 70), body, font=F_SMALL, fill=hexrgb(MID))
    d.text((800, 840), "Only active, in-date prices count. The service must be offered at that clinic.", font=F_BOLD, fill=hexrgb(MID), anchor="mm")
    return save(img, "figure-10-pricing-precedence.png")


def diagram_test_map():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 40), "What every access test must vary", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    center = (800, 450)
    nodes = [
        (80, 120, "Identity", "one person, one account"),
        (615, 95, "Strong sign-in", "required roles and 5-minute rule"),
        (1150, 120, "Membership", "one lifecycle per organization"),
        (1150, 600, "Clinic", "active clinic and clinic link"),
        (615, 690, "Eligibility", "level, credential, supervision"),
        (80, 600, "Isolation", "other organizations unchanged"),
    ]
    for x, y, _t, _b in nodes:
        d.line([(x + 185, y + 75), center], fill=hexrgb(BLUE), width=4)
    for x, y, title, body in nodes:
        rounded(d, (x, y, x + 370, y + 150), PALE, TEAL, 3)
        d.text((x + 185, y + 30), title, font=F_H2, fill=hexrgb(INK), anchor="ma")
        d.text((x + 185, y + 87), body, font=F_SMALL, fill=hexrgb(MID), anchor="ma")
    d.ellipse((650, 300, 950, 600), fill=hexrgb(NAVY))
    d.multiline_text(center, "ROLE\nSCOPE\nCHECK", font=F_H2, fill="white", anchor="mm", spacing=8, align="center")
    return save(img, "figure-12-tester-map.png")


def diagram_example():
    img = canvas(); d = ImageDraw.Draw(img)
    d.text((800, 40), "Worked example: what Dr Sara can do where", font=F_TITLE, fill=hexrgb(INK), anchor="ma")
    rounded(d, (60, 115, 1540, 235), GREEN_FILL, GREEN, 4)
    d.text((800, 145), "One sign-in across three independent provider organizations", font=F_H1, fill=hexrgb(INK), anchor="ma")
    d.text((800, 200), "Each action is checked against organization, clinic, role, status and professional eligibility", font=F_SMALL, fill=hexrgb(INK), anchor="ma")
    cols = [
        (70, "Al Noor", "Owner of the organization\nConsultant at Dubai and Sharjah", "Can: manage members and all clinics; consult at Dubai and Sharjah\nCannot: approve her own credentials; give herself new roles", BLUE),
        (560, "Gulf Heart", "Practice Manager at Gulf Heart Main", "Can: run Gulf Heart Main operations\nCannot: act at Gulf Heart West; approve prices she wrote", TEAL),
        (1050, "Crescent Medical", "Associate Doctor at Crescent Rehab", "Can: work under Dr Lina Mansour's supervision\nCannot: act as a Consultant at Crescent", GOLD),
    ]
    for x, title, role, result, color in cols:
        rounded(d, (x, 280, x + 450, 760), PALE, color, 4)
        d.text((x + 225, 305), title, font=F_H1, fill=hexrgb(color), anchor="ma")
        draw_wrapped(d, (x + 225, 370), role, F_H2, hexrgb(INK), 30, 9, "ma")
        draw_wrapped(d, (x + 225, 520), result, F_BODY, hexrgb(INK), 34, 9, "ma")
    d.text((800, 830), "Removing one role or clinic link changes only that one thing", font=F_BOLD, fill=hexrgb(RED), anchor="mm")
    return save(img, "figure-11-worked-example.png")


# --------------------------------------------------------------------------- document helpers

def add_callout(doc, title, lines, fill=AMBER_FILL, border=GOLD):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, border, "10")
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 140, 180, 140, 180)
    set_cell_width(cell, 7.06)
    p = cell.paragraphs[0]
    style_paragraph(p, after=4, line=1.05)
    font_run(p.add_run(title), 10.5, True, INK)
    for line in lines:
        q = cell.add_paragraph()
        style_paragraph(q, after=3, line=1.08)
        font_run(q.add_run(line), 9.8, False, INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_label(doc, label, text, color=NAVY):
    p = doc.add_paragraph()
    style_paragraph(p, after=6, line=1.12)
    font_run(p.add_run(label + "  "), 10, True, color)
    font_run(p.add_run(text), 10.3)
    return p


def add_role(doc, role, purpose, scope, can_do, cannot_do, granted, removed, mfa, handoff, test_focus):
    add_heading(doc, role, 2)
    intro = add_body(doc, purpose)
    intro.paragraph_format.keep_with_next = True
    add_table(doc, ["Role fact", "Plain-language meaning"], [
        ("Where it applies", scope),
        ("May do", can_do),
        ("Must not do", cannot_do),
        ("Who grants it", granted),
        ("Who removes it", removed),
        ("Strong sign-in", mfa),
        ("Typical handoff", handoff),
        ("Tester focus", test_focus),
    ], [1.5, 5.6], 9.0, TEAL)


STAFF_GRANT_PRIMARY = "A System Administrator, as this person's job, in the staff invitation or through a job change. It cannot be added as an extra responsibility."
STAFF_GRANT_EITHER = "A System Administrator, either as the person's job or as an additional responsibility for someone who already has another compatible job."
STAFF_REMOVE = "A System Administrator through a job change or by removing the responsibility, or through offboarding, which the person's manager or a System Administrator starts. A job change is refused while the person still owns work that the new job cannot handle."
COVERAGE_NOTE = " The organizations and clinics served are RehletShifaa coverage, also granted and ended only by a System Administrator."
STAFF_MFA = "Required on every request. Before first use the person enrols an authenticator."


# --------------------------------------------------------------------------- document

def build_doc(fig):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    for name, size, before, after in [("Heading 1", 18, 10, 5), ("Heading 2", 13.5, 8, 4), ("Heading 3", 11.5, 7, 3)]:
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string("000000")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for sec in doc.sections:
        hp = sec.header.paragraphs[0]
        font_run(hp.add_run("RehletShifaa Users, Roles and Access Guide  |  Edition 2"), 8.5, False, MID)
        add_page_number(sec.footer.paragraphs[0])

    # ---- Cover
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(12)
    p.add_run("RehletShifaa Users, Roles and Access Guide")
    ppr = p._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)
    p2 = doc.add_paragraph()
    style_paragraph(p2, after=12)
    font_run(p2.add_run("A plain-language guide for platform owners, provider organizations, clinic managers, RehletShifaa staff and testers"), 15, False, MID)
    add_body(doc, "Edition 2  |  25 September 2026  |  Replaces Edition 1 of the same date")
    add_callout(doc, "Status: planned, not yet available", [
        "This guide describes how access will work in the first production release of RehletShifaa. That release is not built yet.",
        "Today's development system still works the old way in several important places. Section 2 lists the differences.",
        "Testers: the handbook in Section 12 is the acceptance test for the first production release. Many of its tests fail on today's system by design.",
    ], AMBER_FILL, GOLD)
    add_picture(doc, fig["populations"], "Four user populations: platform governance, RehletShifaa teams, provider organizations and patients.", 6.4)
    p = doc.add_paragraph()
    style_paragraph(p, before=8, after=0)
    font_run(p.add_run("Document owner: RehletShifaa.  Source: Customer Governance Target Architecture, revision 3 (final), decisions D1 to D35."), 9, True, NAVY)
    add_page_break(doc)

    # ---- 1 How to use
    add_heading(doc, "1 How to use this guide", 1)
    add_body(doc, "The model is simple. A role says what a person may do. A scope says where the role applies. A person gets access only when everything that applies is active at the same time: the person, their membership or coverage, the organization, the clinic, the role, and any professional qualification the work needs.")
    add_heading(doc, "Reader paths", 2)
    add_table(doc, ["Reader", "Start here", "Then read"], [
        ("Provider owner or buyer", "Sections 3, 6 and 9", "Worked example (Section 13)"),
        ("Clinic manager", "Section 6 (organization and clinics)", "Sections 8, 9 and 10"),
        ("RehletShifaa manager", "Sections 4 and 5", "Section 10 (joining and leaving)"),
        ("Tester", "Section 2, then Section 12", "Sections 8 to 11 for expected results"),
        ("Implementation team", "Section 2 and Section 14", "All role profiles"),
    ], [1.7, 2.4, 3.0], 9.3)
    add_heading(doc, "How firm is each statement?", 2)
    add_body(doc, "Almost everything in this guide is an agreed rule. Where a statement rests on a sensible default rather than an explicit business decision, or has not been decided at all, it is labelled:")
    add_table(doc, ["Label", "Meaning"], [
        ("Agreed rule", "Decided in the architecture. Unlabelled statements are agreed rules."),
        ("Adopted default", "The architecture's recommended default. It stands unless the business explicitly changes it."),
        ("Not yet decided", "The architecture does not answer this. Do not build or test an answer until it is decided. See Section 14."),
        ("Names pending", "English and Arabic role and clinic names are still in native-language review. Wording may change; meaning will not."),
    ], [1.6, 5.5], 9.3)
    add_heading(doc, "Contents", 2)
    add_numbered(doc, [
        "How to use this guide",
        "What exists today and what is planned",
        "The model in one page",
        "Platform governance roles",
        "RehletShifaa internal staff roles",
        "Provider organizations, clinics and roles",
        "Patients and patient representatives",
        "Scopes, role combinations and professional eligibility",
        "One person in several organizations: signing in and switching",
        "Joining, activation, suspension and leaving",
        "Services, prices, schedules and cases",
        "Tester handbook",
        "Worked example",
        "Quick reference, glossary and decision status",
    ])

    # ---- 2 Today vs planned
    add_heading(doc, "2 What exists today and what is planned", 1)
    add_body(doc, "The first production release (called P1 in the architecture) is delivered as one package. None of its parts goes live on its own, and there is no rollback to the old behaviour once it is live. Until P1 is built, the development system behaves as shown in the middle column.")
    add_table(doc, ["Area", "Today's development system", "First production release"], [
        ("Administrator and auditor access to medical data", "The old administrator and auditor roles can open the medical documents of any case.", "System Administrator and Compliance Auditor have no access to patient cases or documents."),
        ("Strong sign-in (second factor)", "Not configured. Password only.", "Required for the owner, all internal staff, provider managers and clinicians."),
        ("Clinics and branches", "Do not exist. The organization is the only unit; locations are free text.", "Each organization has clinics with their own status, services, prices and schedules."),
        ("Inviting a person who already has an account", "The person is attached to the organization silently, by email.", "The person receives an offer and must accept. Nothing is active before acceptance."),
        ("Consultant or Associate Doctor level", "One platform-wide type, sometimes a default value rather than a real declaration.", "Declared level plus a separate credential decision for each organization."),
        ("Internal staff jobs", "Granted as sign-in system roles.", "Granted as RehletShifaa assignments: one job plus optional extra responsibilities."),
        ("Platform owner and administrators", "A setup account holds broad legacy roles.", "An owner relationship plus approved administrator appointments."),
        ("Staff joining and leaving", "Separate manual commands.", "Guided invitation, activation at first strong sign-in, and orchestrated offboarding."),
        ("RehletShifaa coverage of a provider", "A provider owner can currently remove RehletShifaa's credential reviewer.", "Providers cannot see, change or remove RehletShifaa coverage."),
    ], [1.75, 2.6, 2.75], 8.8)
    add_body(doc, "Evidence: architecture review sections 2 and 14A.1, checked against the code at commit e4c7552 on 25 September 2026.", italic=True)
    add_callout(doc, "What this means for testers", [
        "Run the Section 12 handbook only against a build that contains the whole first production release.",
        "On today's system, a failed test in Section 12 is expected. Record it as \"not yet built\", not as a defect.",
    ], AMBER_FILL, GOLD)

    # ---- 3 Model
    add_heading(doc, "3 The model in one page", 1)
    add_heading(doc, "Seven facts every reader should know", 2)
    add_numbered(doc, [
        "One human has one sign-in identity across the whole platform.",
        "A provider organization is the legal or contracting party. It is the boundary between providers.",
        "A clinic, branch, hospital or virtual-care site is a facility inside exactly one provider organization. The screens call these \"Clinics & branches\".",
        "A provider professional may belong to several organizations and work at several clinics in each one.",
        "The same person may hold different compatible roles in different organizations or clinics.",
        "Ending one role, clinic link or membership never switches off access anywhere else.",
        "Switching off a person's sign-in completely is reserved for account compromise or an authorized platform-wide action.",
    ])
    add_heading(doc, "The provider hierarchy", 2)
    add_picture(doc, fig["provider_tree"], "Dr Sara has one identity with separate memberships, roles and clinic links in three organizations.", 6.8)
    add_caption(doc, "Figure 1  One person can work across organizations and clinics")
    add_bullets(doc, [
        "The sign-in identity and the practitioner profile are shared across the platform.",
        "Each organization membership starts and ends on its own.",
        "Each clinic link and each role names exactly where it applies.",
    ])
    add_heading(doc, "Two ways to reach a provider organization", 2)
    add_body(doc, "Provider staff reach their organization through a provider membership. RehletShifaa employees who serve that organization, for example a Credentialing Specialist, reach it through RehletShifaa coverage. The two never mix.")
    add_picture(doc, fig["two_paths"], "Provider users reach an organization through membership, clinic links and roles; RehletShifaa staff through centrally granted coverage.", 6.8)
    add_caption(doc, "Figure 2  Provider membership and RehletShifaa coverage are separate")
    add_table(doc, ["", "Provider membership", "RehletShifaa coverage"], [
        ("Who has it", "Organization Owners, Practice Managers, Consultants, Associate Doctors, Consultant Assistants", "RehletShifaa staff such as Provider Operations, Credentialing Specialists and Care Coordination"),
        ("How it starts", "The person accepts an offer", "A System Administrator grants it"),
        ("Needs a clinic link", "Yes, for clinic-specific roles", "No"),
        ("Visible to the provider", "Yes, in the organization's people lists", "No. Shown only in RehletShifaa's Control Center"),
        ("Who can end it", "The organization's authorized managers", "Only a RehletShifaa System Administrator"),
        ("Can one person have both?", "No. A RehletShifaa employee cannot also be a provider member (adopted default).", "No, for the same reason."),
    ], [1.6, 2.75, 2.75], 8.9)
    add_page_break(doc)

    # ---- 4 Governance
    add_heading(doc, "4 Platform governance roles", 1)
    add_body(doc, "These roles govern the RehletShifaa platform itself. They are separate from daily care work and from owning a provider organization.")
    add_picture(doc, fig["governance"], "The Platform Account Owner approves administrator changes; System Administrators manage staff and access; neither has clinical access.", 6.8)
    add_caption(doc, "Figure 3  Governance and daily work are separate")
    add_role(doc, "Platform Account Owner",
             "The contractual root of trust for the RehletShifaa platform. Ownership is a protected relationship, not a role that anyone can grant.",
             "Platform governance actions only.",
             "Approve or reject System Administrator appointments and removals; transfer ownership; authorize the first administrators at handover; restore administration when no administrator is left; remove an administrator in an emergency; cancel a vendor recovery; read the governance summary.",
             "Read patient cases or documents; run finance or provider operations; grant ordinary staff jobs; act as the owner of any provider organization because of platform ownership.",
             "Not granted like a role. The first owner is named at installation and claims ownership with strong sign-in. After that, only the current owner can nominate a successor, and the successor must accept.",
             "Ends only when a nominated successor accepts, or when a vendor recovery completes. A System Administrator can never transfer or remove ownership.",
             "Required for every owner action. Sensitive actions also need a sign-in within the last five minutes.",
             "Approves administrator changes. System Administrators do the day-to-day staff and access work.",
             "Ownership alone gives no clinical, operational, financial, provider or patient access. Transfer and emergency actions are recorded.")
    add_heading(doc, "Transferring platform ownership", 2)
    add_picture(doc, fig["ownership"], "Four-step ownership transfer: nominate, set up strong sign-in, accept, ownership moves.", 6.8)
    add_caption(doc, "Figure 4  Ownership moves only when the successor accepts")
    add_body(doc, "Vendor recovery (adopted default: 72-hour waiting period) is for the rare case where the owner cannot be reached and no administrator remains. It is verified outside the platform, everyone is notified, and it can be cancelled during the waiting period.")
    add_role(doc, "System Administrator",
             "Manages RehletShifaa staff and access without becoming a clinical superuser. Two active holders are recommended.",
             "Platform access governance and staff administration.",
             "Invite staff; change a person's job; add or remove compatible extra responsibilities; grant and end RehletShifaa coverage of provider organizations; view effective access and the access audit; request or approve another administrator's appointment or removal.",
             "Approve their own request; grant or remove their own access; open patient cases or documents; make finance, credential, patient identity, provider or care journey decisions; bypass role conflicts; remove the last administrator; give someone a job through the People page instead of a job change.",
             "An administrator requests the appointment and another administrator or the owner approves it (see the table below). The person must already be an active RehletShifaa staff member.",
             "A removal request approved by another administrator or the owner; the person's own resignation; or the owner's emergency removal. The last administrator cannot be removed except by the owner's emergency removal.",
             "Required on every request. Administrator changes, access grants and staff invitations also need a sign-in within the last five minutes.",
             "Hands operational work to the right staff job. Another administrator or the owner approves administrator changes.",
             "No self-approval, no clinical access, last-administrator protection, and owner recovery when no administrator remains.")
    add_heading(doc, "Administrator appointment and removal rules", 2)
    add_table(doc, ["Situation", "Who starts it", "Who approves it", "Protection"], [
        ("Handover at installation", "The vendor's installation file names the first administrators", "The owner, when claiming ownership", "The vendor keeps no account after handover"),
        ("Two or more administrators", "Any administrator", "Another administrator (not the person affected) or the owner", "Requester, approver and person affected must differ"),
        ("Exactly one administrator", "That administrator", "The owner", "The last administrator stays until the replacement is active"),
        ("No administrator after handover", "The owner (\"Restore administration\")", "No second approval", "Strong sign-in, recent sign-in, reason and a critical notice"),
        ("Owner unreachable and no administrator", "Vendor recovery", "Waiting period, can be cancelled", "Verified outside the platform"),
        ("Normal removal", "Any administrator", "Another administrator or the owner", "Cannot remove the last administrator"),
        ("Resignation", "The administrator", "None", "Cannot remove the last administrator"),
        ("Emergency removal", "The owner", "None", "Explicit warning; may leave zero administrators until restored"),
        ("Owner also becomes an administrator", "Another administrator", "Another administrator", "Never self-requested or self-approved (adopted default: allowed)"),
    ], [1.6, 1.65, 1.75, 2.1], 8.5)
    add_body(doc, "A request that nobody approves expires after 72 hours (adopted default). The screens say: \"Add the replacement first, then remove.\"")

    # ---- 5 Internal staff
    add_heading(doc, "5 RehletShifaa internal staff roles", 1)
    add_body(doc, "Internal staff are RehletShifaa employees. They stay RehletShifaa employees even when they serve a provider organization, never appear in a provider's people lists, and cannot be invited into a provider organization.")
    add_heading(doc, "How internal jobs work", 2)
    add_bullets(doc, [
        "Each staff member has exactly one job. The only exception is a person who is a System Administrator and nothing else (\"administration only\").",
        "Operational jobs (coordination, travel, finance and provider operations) can only be a person's main job, because they carry work queues and reporting lines.",
        "Five independent functions can be a job or an extra responsibility: Credentialing Specialist, Patient Identity Reviewer, Care Journey Manager, Care Journey Approver and Compliance Auditor. System Administrator can also be added as a protected extra responsibility.",
        "Some combinations are always refused for one person: Credentialing Specialist with Provider Operations; Care Journey Manager with Care Journey Approver; Compliance Auditor with System Administrator.",
        "The job title on screen is for display. Access comes from the person's assignments and their scope.",
        "A new staff member has no access until their first sign-in with strong sign-in.",
    ])
    staff = [
        ("Care Coordination Manager", "Leads the coordination team and supervises patient cases for the provider organizations the team serves.",
         "The whole platform for coordination work, limited to the provider organizations the manager serves.",
         "Supervise coordination work; view team queues; transfer cases; set up coordination pools and routing rules for the organizations served; do coordination casework.",
         "Act for organizations not served; do finance, travel, credentialing, provider administration or access administration.",
         STAFF_GRANT_PRIMARY + COVERAGE_NOTE, STAFF_REMOVE,
         "Receives staff setup from a System Administrator; hands cases to Care Coordinators and other teams.",
         "Team visibility, transfer and organization boundaries. Promoting a Care Coordinator to this job is not blocked by the cases they own."),
        ("Care Coordinator", "Owns and moves forward assigned coordination cases, connecting patients, providers and internal teams.",
         "Coordination work for the provider organizations served. Which cases they can open depends on case ownership and routing.",
         "Take ownership of eligible cases; work assigned coordination tasks; hand work to the next team as the workflow requires.",
         "Supervise the whole team; configure pools; open unrelated cases; do travel, finance, credentialing or administration work.",
         STAFF_GRANT_PRIMARY + COVERAGE_NOTE, STAFF_REMOVE,
         "Escalates or transfers to the Care Coordination Manager; hands approved work to travel, finance or provider teams.",
         "Access to owned and assigned cases; denial for unrelated cases and organizations."),
        ("Travel and Logistics Manager", "Supervises travel and logistics work for patient journeys.",
         "Travel casework across the platform, with supervision.",
         "Supervise travel cases; review team work; see the team's travel work; manage handoffs within travel.",
         "Change clinical decisions; set finance policy; manage provider membership; open case data beyond what travel work needs.",
         STAFF_GRANT_PRIMARY, STAFF_REMOVE,
         "Receives travel needs from coordination and hands completed arrangements back.",
         "Supervisory view limited to travel work."),
        ("Travel and Logistics Specialist", "Plans and records travel arrangements for assigned cases.",
         "Assigned travel cases.",
         "Build travel plans; record arrival and logistics milestones; complete assigned operations steps.",
         "Supervise the travel team; open unrelated cases; approve finance or clinical decisions.",
         STAFF_GRANT_PRIMARY, STAFF_REMOVE,
         "Works from coordination requirements and passes completed arrangements on.",
         "Access limited to assigned travel cases."),
        ("Commercial and Finance Manager", "Owns commercial policy and supervises the financial side of patient cases.",
         "Finance casework and commercial policy across the platform.",
         "Set margin and deposit policy; manage exchange rates and service templates; supervise finance casework; do finance work when needed.",
         "Open unrelated clinical material; grant access; make credential decisions.",
         STAFF_GRANT_PRIMARY, STAFF_REMOVE,
         "Sets the policy that Finance Officers apply and handles exceptions.",
         "Manager-only policy, exchange-rate and template actions."),
        ("Finance Officer", "Carries out the financial steps of assigned cases under the approved commercial policy.",
         "Assigned finance cases.",
         "Record deposits and payments; approve commercial terms for assigned cases; complete assigned finance tasks.",
         "Change margin, deposit, exchange-rate or template policy; open unrelated cases; administer users.",
         STAFF_GRANT_PRIMARY, STAFF_REMOVE,
         "Receives approved terms and hands financial completion back to the patient journey.",
         "Assigned-case limits; denial for policy changes."),
        ("Provider Operations Manager", "Leads the setup and operational governance of provider organizations.",
         "Creating providers across the platform, plus the organizations or clinics this person covers.",
         "Create provider organizations with their first clinic; activate or suspend providers; manage provider and clinic setup; send membership offers for covered organizations; manage direct consultant setup.",
         "Decide credentials for a clinician they invited or manage; become a provider employee; open patient cases.",
         STAFF_GRANT_PRIMARY + COVERAGE_NOTE, STAFF_REMOVE,
         "Prepares providers, then hands credential decisions to an independent Credentialing Specialist.",
         "Coverage boundaries; independence from credentialing."),
        ("Provider Operations Specialist", "Sets up providers and clinicians for covered organizations under the manager's oversight.",
         "Covered organizations or clinics.",
         "View and update covered providers; send membership offers; perform approved direct setup tasks.",
         "Activate or suspend providers; decide credentials for clinicians they invited or manage; act outside coverage.",
         STAFF_GRANT_PRIMARY + COVERAGE_NOTE, STAFF_REMOVE,
         "Prepares the organization and clinicians for independent credential review and manager activation.",
         "No activation authority; no credential decisions for their own invitees."),
        ("Credentialing Specialist", "Makes the independent credential decision for clinicians, separately for each provider organization.",
         "Covered organizations or clinics.",
         "View credential files; review evidence; ask for more information; verify, reject or suspend credentials, including the level (Consultant or Associate Doctor) for that organization; decide direct clinician credentials.",
         "Review their own credentials; decide on a clinician they invited or manage; hold a Provider Operations job at the same time.",
         STAFF_GRANT_EITHER + COVERAGE_NOTE, STAFF_REMOVE,
         "Receives a complete file from provider operations or the clinician and returns an independent decision.",
         "Self-review denial, inviter and reviewer separation, and clinic-level coverage."),
        ("Patient Identity Reviewer", "Reviews patients' legal identity evidence and records a decision.",
         "The patient identity review queue.",
         "Open the identity review queue; inspect the identity evidence; approve, reject or ask for more information.",
         "Read clinical documents or unrelated case details; administer staff; make credential decisions.",
         STAFF_GRANT_EITHER, STAFF_REMOVE,
         "Returns the identity decision to the patient journey.",
         "Sees identity evidence only, never general clinical documents."),
        ("Care Journey Manager", "Designs and prepares reusable care journey templates.",
         "Care journey design.",
         "Create and edit draft journeys; validate and simulate them; submit them for approval.",
         "Approve, publish or retire their own journey; also hold Care Journey Approver.",
         STAFF_GRANT_EITHER, STAFF_REMOVE,
         "Submits a finished draft to a Care Journey Approver.",
         "Author and approver separation; no patient case access."),
        ("Care Journey Approver", "Independently approves and publishes care journey templates.",
         "Care journey approval.",
         "Review submitted journeys; approve, publish or retire journey versions.",
         "Author the journey they approve; also hold Care Journey Manager; open patient cases.",
         STAFF_GRANT_EITHER, STAFF_REMOVE,
         "Receives a submitted journey and publishes or rejects it.",
         "Maker and checker separation; no patient case access."),
        ("Compliance Auditor", "Independently reviews roles, assignments and who can do what.",
         "Access audit across the platform.",
         "View role definitions; view effective access; review the access audit history.",
         "Open patient cases or documents; administer staff; change roles; also be a System Administrator. Case-level clinical audit is a future capability, not part of this release.",
         STAFF_GRANT_EITHER, STAFF_REMOVE,
         "Reviews evidence produced by access governance without changing it.",
         "Read-only audit access; no clinical access; conflict with System Administrator."),
    ]
    for role, purpose, scope, can, cannot, grant, remove, handoff, focus in staff:
        add_role(doc, role, purpose, scope, can, cannot, grant, remove, STAFF_MFA, handoff, focus)

    # ---- 6 Provider
    add_heading(doc, "6 Provider organizations, clinics and roles", 1)
    add_body(doc, "A provider organization is the legal or contracting party. A facility is one operating site inside that organization; the screens call facilities \"Clinics & branches\". Each clinic belongs to exactly one organization.")
    add_heading(doc, "One organization with several clinics", 2)
    add_picture(doc, fig["multi_clinic"], "Al Noor with four clinics; bars show which people reach which clinics.", 6.8)
    add_caption(doc, "Figure 5  Organization-wide roles reach every clinic; clinic-specific roles reach only the named clinics")
    add_bullets(doc, [
        "An organization can have any number of clinics, hospitals, medical centers, diagnostic centers, rehabilitation centers and virtual-care sites.",
        "A site that is a separate legal or contracting party is its own provider organization, not a branch.",
        "A clinic is Onboarding, Active, Suspended or Closed. A physical clinic cannot become active without an address.",
        "An organization is created together with its first clinic, which becomes its default clinic.",
        "An organization cannot be active without at least one active clinic.",
        "Exactly one clinic is the default. It pre-fills forms and anchors older records. It never grants access by itself.",
        "Closing the default clinic requires choosing a replacement; otherwise the organization shows \"default clinic required\". Suspending the default clinic does not move the default.",
        "Suspending one clinic stops work at that clinic only. Sibling clinics carry on.",
    ])
    add_callout(doc, "Existing providers after the upgrade", [
        "Each existing organization receives one default clinic, created in the Onboarding state because today's records hold no address.",
        "The organization owner or RehletShifaa Provider Operations must add the address (and confirm the clinic type where flagged) before that clinic can become active.",
        "Existing staff and clinicians are linked to the default clinic. Organization-wide roles keep working unchanged.",
    ], LIGHT, BLUE)
    add_heading(doc, "Organization-wide versus clinic-specific", 2)
    add_table(doc, ["Action", "Organization-wide manager", "Manager of Dubai and Sharjah only"], [
        ("Work at Dubai or Sharjah", "Yes", "Yes"),
        ("Work at Abu Dhabi Hospital", "Yes", "No"),
        ("Work at a clinic opened next month", "Yes, automatically", "No, until the clinic is added to the role"),
        ("Change organization defaults (for example default prices)", "Yes, if the role allows it", "No. Organization-level settings need an organization-wide role"),
        ("Manage clinicians", "Clinicians anywhere in the organization", "Only managed clinicians at Dubai and Sharjah"),
        ("Invite new members", "Only if member invitation has been granted", "Only if member invitation has been granted"),
    ], [2.6, 2.25, 2.25], 9.0)
    provider_roles = [
        ("Organization Owner", "Owns the provider organization and governs its members and operations. This is a different role from the Platform Account Owner.",
         "The whole organization, including every current and future clinic.",
         "Manage provider members and organization-level settings; send membership offers; manage clinics and branches; hold other compatible provider roles where the rules allow.",
         "Grant platform or RehletShifaa staff access; see, change or remove RehletShifaa coverage; bypass self-approval rules; give themselves a role; see another organization; do clinical work without the clinical qualification and credential checks.",
         "Given through a membership offer that includes this role and that the person accepts. Offers can come from RehletShifaa Provider Operations for an organization it covers, or from an authorized manager of the organization. Who may grant this role to an existing member is not yet decided (Section 14).",
         "Not yet decided: who may remove an Organization Owner, whether an organization must always keep at least one, and how ownership passes to someone else (Section 14). Ending the person's membership ends the role.",
         "Required at launch, on every request."),
        ("Practice Manager", "Runs provider operations for one, several or all clinics of the organization.",
         "Organization-wide, or only the listed clinics. There is no separate Branch Manager role.",
         "Manage operations and managed clinicians within scope; manage schedules, services and prices where the role allows; invite members only if member invitation has been granted separately.",
         "Act at an unlisted sibling clinic; change organization-level settings when clinic-limited; grant platform access; approve a price they wrote; bypass qualification or clinic-link rules.",
         "Offer acceptance, or a role grant by an authorized provider manager (normally the Organization Owner). Nobody can grant a role to themselves.",
         "The same authorized provider managers. A clinic-limited role loses a clinic when the person's link to that clinic ends, and ends if no clinic remains.",
         "Required at launch, on every request."),
        ("Consultant", "Provides consultant-level services in the organization and clinics named in the role.",
         "Listed clinics where the person has an active clinic link, or organization-wide when granted that way; plus their own professional workspace.",
         "Work as a consultant at those clinics; manage their own schedule and eligible services; supervise Associate Doctors through an active supervision relationship.",
         "Work as a consultant without a Consultant declaration and a Consultant-level credential decision for this organization; work at a clinic without an active clinic link; approve their own credentials or a price they wrote.",
         "Offer acceptance or a grant by an authorized provider manager, and only when the checks in Section 8 pass.",
         "The same provider managers; ending the clinic link or membership; or the credential being suspended.",
         "Required at launch (confirmed business decision)."),
        ("Associate Doctor", "Provides clinical services under the supervision of an eligible Consultant.",
         "Listed clinics with an active clinic link, plus their own workspace and the supervision relationship.",
         "Work as an Associate Doctor at those clinics under an active supervisor; maintain their own workspace.",
         "Work without an active supervising Consultant in the same organization; act as a Consultant there; act at another clinic or organization without the matching role and clinic link.",
         "Offer acceptance or a grant by an authorized provider manager, only with an Associate or Consultant credential decision for this organization and an active supervisor.",
         "The same provider managers; ending the clinic link or membership. If supervision ends, the role cannot be used until a supervisor is in place again.",
         "Required at launch (confirmed business decision)."),
        ("Consultant Assistant", "Supports selected clinicians and clinic operations without becoming a clinician.",
         "Listed clinics, and the clinicians the assistant is explicitly assigned to support.",
         "Support the assigned clinicians and clinics; perform allowed administrative tasks within that scope.",
         "Hold Consultant or Associate Doctor at the same clinic; do clinical work through this role; act for clinicians or clinics not assigned.",
         "Offer acceptance or a grant by an authorized provider manager.",
         "The same provider managers; ending the clinic link or membership.",
         "Not required by this role itself. Required if the same person holds any role that needs it, anywhere on the platform."),
    ]
    for role, purpose, scope, can, cannot, grant, remove, mfa in provider_roles:
        focus = {
            "Organization Owner": "Reach to every clinic including new ones; denial of every platform, staff and coverage action; no self-granting.",
            "Practice Manager": "Test both variants. A clinic-limited manager is denied at sibling clinics and on organization defaults.",
            "Consultant": "Declared level, organization credential level, active clinic link and role scope must all be present.",
            "Associate Doctor": "Active supervision is mandatory; a Consultant-level credential also satisfies the Associate requirement.",
            "Consultant Assistant": "Same-clinic conflict with clinical roles; denial outside assigned clinicians and clinics.",
        }[role]
        handoff = {
            "Organization Owner": "Delegates daily work to Practice Managers and clinical work to eligible clinicians.",
            "Practice Manager": "Coordinates staff, clinicians and clinic operations within scope.",
            "Consultant": "Credential decisions come from an independent RehletShifaa Credentialing Specialist.",
            "Associate Doctor": "Works under the supervising Consultant and the organization's credential decision.",
            "Consultant Assistant": "Escalates clinical decisions to the responsible clinician.",
        }[role]
        add_role(doc, role, purpose, scope, can, cannot, grant, remove, mfa, handoff, focus)
    add_callout(doc, "Clinical work in the first release", [
        "Provider clinicians see the case summary made available for cases assigned to them in their organization and clinic. Wider clinical capabilities (clinical records, appointments) are future work.",
        "Wherever this guide says a Consultant may \"work\" or \"consult\", it means the provider capabilities that exist in that release, always within these rules.",
    ], LIGHT, BLUE)

    # ---- 7 Patients
    add_heading(doc, "7 Patients and patient representatives", 1)
    add_body(doc, "The patient experience is not changed by this model. These roles are listed so that readers can see how they relate to the others.")
    add_role(doc, "Patient", "Uses the patient portal for their own identity, care requests, documents, messages and journey steps.",
             "Their own care.",
             "Use the patient portal; keep permitted profile details up to date; see and act on their own care; keep patient access after joining a provider organization.",
             "Receive provider, staff or platform authority from the Patient role.",
             "Created through patient sign-up and account activation.",
             "Unchanged by this model.",
             "Password sign-in is enough for patient use. If the same person later takes on a role that needs strong sign-in (for example Organization Owner), strong sign-in then applies to all their sign-ins.",
             "Works with care coordination, providers, travel and finance through the patient journey.",
             "Joining a provider organization does not remove Patient access; provider authority does not start before acceptance.")
    add_role(doc, "Patient Representative", "Acts for a patient in the patient portal, for example a relative who submits the case.",
             "The represented patient's cases, where the representative's authorization is recorded.",
             "Use the portal for the represented patient's permitted actions.",
             "Act for unrelated patients; gain provider, staff or platform authority from this role.",
             "Existing patient-portal representative flow. Relationship-based representative grants are future work.",
             "Unchanged by this model.",
             "Same as Patient.",
             "Communicates with the care team within the recorded authorization.",
             "Existing behaviour only (regression test): no access to other patients' information.")
    add_heading(doc, "Things that are not business roles", 2)
    add_table(doc, ["Item", "What it is", "Why it is not a role"], [
        ("\"Strong sign-in required\" marker", "A sign-in setting that forces a second factor", "It never grants any business permission"),
        ("Vendor operator", "The installer used during setup or recovery", "Not a platform user; it leaves after handover"),
        ("\"Provider user\"", "A group label", "Access comes from Organization Owner, Practice Manager, Consultant, Associate Doctor or Consultant Assistant roles"),
        ("\"Internal staff\"", "A group label", "Access comes from the person's staff job, extra responsibilities and coverage"),
    ], [1.6, 2.4, 3.1], 9.0)
    add_heading(doc, "Not in the launch catalogue", 2)
    add_bullets(doc, [
        "Support Specialist and Support Manager: deferred.",
        "Credentialing Manager: not a separate role.",
        "Journey Specialist roles: not included.",
        "Platform Administrator, Access Governance Manager and Support Agent: retired.",
        "Branch Manager: not a role. A Practice Manager limited to certain clinics covers it.",
        "The old sign-in system staff roles: removed before go-live; they never grant access in production.",
    ])

    # ---- 8 Scopes
    add_heading(doc, "8 Scopes, role combinations and professional eligibility", 1)
    add_heading(doc, "Where a role applies", 2)
    add_body(doc, "A role without the right scope cannot reach the record. The platform works out the organization and clinic from the stored record, never from what the browser sends.")
    add_table(doc, ["Scope", "Plain meaning", "Used by"], [
        ("Organization", "The whole provider organization, including clinics added later", "Organization Owner; organization-wide Practice Manager; clinicians granted organization-wide"),
        ("Selected clinics (provider)", "Only the listed clinics of one organization, where the person has an active clinic link", "Clinic-limited Practice Manager, Consultant, Associate Doctor, Consultant Assistant"),
        ("Managed clinicians", "Only clinicians linked through an active management relationship", "Practice Managers and support roles"),
        ("Self", "Only the person's own profile and professional records", "Every clinician and assistant"),
        ("Platform", "A named capability across RehletShifaa", "System Administrator and platform-wide staff jobs"),
        ("Covered organizations", "Only the provider organizations assigned to a RehletShifaa staff member", "Provider Operations, Credentialing, Care Coordination"),
        ("Covered clinics", "Only listed clinics of a covered organization. No clinic link needed", "Provider Operations and Credentialing when limited to certain clinics"),
    ], [1.6, 3.0, 2.5], 8.9)
    add_body(doc, "Which individual cases a coordinator, travel or finance staff member can open is decided by case ownership and assignment in the workflow, not by a separate scope.")
    add_heading(doc, "Internal staff combinations", 2)
    add_table(doc, ["Combination", "Result"], [
        ("One job", "Required. Only a person who is a System Administrator and nothing else may have no job."),
        ("Extra responsibilities", "Credentialing Specialist, Patient Identity Reviewer, Care Journey Manager, Care Journey Approver, Compliance Auditor and System Administrator, when compatible (adopted default)."),
        ("Credentialing Specialist with Provider Operations", "Refused. Whoever invites or manages clinicians cannot also decide their credentials."),
        ("Care Journey Manager with Care Journey Approver", "Refused. The author cannot approve their own work."),
        ("Compliance Auditor with System Administrator", "Refused. The auditor cannot audit their own administration."),
        ("Any internal job with a provider role", "Refused. RehletShifaa staff are never provider members (adopted default)."),
    ], [2.4, 4.7], 9.1)
    add_heading(doc, "Provider role combinations", 2)
    add_table(doc, ["Held by one person", "Result"], [
        ("Organization Owner and Practice Manager", "Allowed."),
        ("Organization Owner or Practice Manager with Consultant Assistant", "Allowed."),
        ("Organization Owner or Practice Manager with Consultant", "Allowed when the clinical checks pass. No self-credentialing and no approving prices they wrote."),
        ("Organization Owner or Practice Manager with Associate Doctor", "Allowed when the clinical checks and supervision are in place."),
        ("Consultant and Associate Doctor", "Allowed only at different clinics or in different organizations, for a Consultant-qualified person with an active supervisor for the Associate work (adopted default)."),
        ("Consultant Assistant with Consultant or Associate Doctor", "Refused at the same clinic."),
    ], [2.8, 4.3], 9.1)
    add_heading(doc, "Four facts needed for clinical work", 2)
    add_numbered(doc, [
        "Declared level. The professional states, once for the whole platform, whether they are a Consultant or an Associate Doctor. A level carried over from the old system must be confirmed by the professional before any new Consultant role.",
        "Credential decision for this organization. An independent RehletShifaa Credentialing Specialist verifies the level separately for each organization. The organization cannot verify its own clinicians, and evidence is not shared between organizations automatically.",
        "Local role. The organization grants a Consultant or Associate Doctor role for the organization or for listed clinics.",
        "Clinic link. The person has an active link to every clinic named in the role.",
    ])
    add_body(doc, "No local role can raise the declared or verified level. A Consultant role needs a Consultant declaration and a Consultant-level credential decision in that organization. An Associate Doctor role accepts an Associate or Consultant credential decision and needs an active supervising Consultant in the same organization. A person with no declared level cannot receive any clinical role.")

    # ---- 9 Context switching
    add_heading(doc, "9 One person in several organizations: signing in and switching", 1)
    add_body(doc, "A person who works for several organizations signs in once with one account. In the provider workspace (\"My Practice\") they first choose an organization and then a clinic.")
    add_picture(doc, fig["context"], "Sign in once, choose organization, choose clinic; the server re-checks every action.", 6.8)
    add_caption(doc, "Figure 6  The selector changes the view; the server decides access")
    add_bullets(doc, [
        "The organization list shows only organizations where the person has an active membership. The clinic list shows only clinics they are linked to or reach through an organization-wide role.",
        "Every screen and every action is checked again on the server for that organization and clinic.",
        "A clinician page shows only the engagements the viewer is allowed to see. A Gulf Heart manager looking at Dr Sara sees her Gulf Heart work, not her Al Noor or Crescent work.",
        "Search, notifications and audit views are filtered the same way.",
        "Membership offers appear as a card with Accept and Decline.",
    ])
    add_heading(doc, "Strong sign-in belongs to the person, not to one context", 2)
    add_body(doc, "If any of a person's roles anywhere needs strong sign-in, every sign-in by that person needs it, including when they work in a context that would not need it on its own. The requirement is removed only when the last role that needs it ends.")
    add_table(doc, ["Person", "Strong sign-in?", "Why"], [
        ("Dr Sara (Owner at Al Noor, Associate Doctor at Crescent)", "Always", "Owner and clinician roles both require it"),
        ("A patient who accepts an Organization Owner offer", "Always, from acceptance", "The new role requires it; it also applies to patient use"),
        ("A Consultant Assistant with no other role", "No", "The role does not require it"),
        ("A Practice Manager whose last manager role ends, and who is not a clinician", "No longer", "Removed when no qualifying role remains"),
    ], [2.9, 1.5, 2.7], 9.0)

    # ---- 10 Lifecycle
    add_heading(doc, "10 Joining, activation, suspension and leaving", 1)
    add_heading(doc, "Joining a provider organization: offers", 2)
    add_body(doc, "Typing an email never attaches a person to an organization. The organization sends an offer; the person sees the organization, clinics, proposed roles and dates, and then accepts or declines.")
    add_picture(doc, fig["offer"], "Five-step offer flow from sending to access starting after acceptance.", 6.8)
    add_caption(doc, "Figure 7  Provider access starts only after the person accepts")
    add_table(doc, ["Offer state", "Meaning"], [
        ("Awaiting identity", "The platform is finding or creating the one account for this email."),
        ("Pending", "The person can review and accept or decline. Nothing is active yet."),
        ("Accepted", "Membership, clinic links and roles start together. A practitioner profile is created once, if needed."),
        ("Declined", "This offer ends. Nothing else changes."),
        ("Expired", "Not answered in time. Adopted default: 14 days."),
        ("Cancelled", "The sender withdrew the offer before acceptance."),
        ("Review required", "The email matches more than one account in a way that cannot be resolved safely, or belongs to a RehletShifaa employee. RehletShifaa governance reviews it; by default it is refused."),
    ], [1.6, 5.5], 9.2)
    add_bullets(doc, [
        "The sender always sees \"Invitation sent\", whether the email is new, belongs to an existing account or to a RehletShifaa employee. No organization can use invitations to find out who has an account.",
        "If a proposed role does not fit the person (for example Consultant for someone declared as Associate Doctor), the offer is still sent without that role and the sender sees a neutral \"some roles need review\" message.",
        "A manager role can only be accepted with strong sign-in.",
    ])
    add_heading(doc, "Internal staff activation", 2)
    add_numbered(doc, [
        "A System Administrator creates the invitation with one job, any compatible extra responsibilities, and the operational setup (manager, pools, organizations served).",
        "The invitee verifies their email, sets a password, enrols an authenticator and signs in.",
        "The first request made with strong sign-in activates the staff record, exactly once.",
        "Until then the person has no access at all. A password-only sign-in never activates a staff record.",
    ])
    add_heading(doc, "Suspension and ending: what changes and what stays", 2)
    add_picture(doc, fig["lifecycle"], "Five local and global actions and what each one affects.", 6.8)
    add_caption(doc, "Figure 8  Local actions do not spread")
    add_table(doc, ["Action", "Who does it", "What ends", "What stays"], [
        ("End one clinic link", "Authorized provider manager", "That clinic from the person's clinic-limited roles; a role left with no clinic ends", "Other clinics, organizations, the membership and RehletShifaa coverage"),
        ("End one provider membership", "Authorized provider manager", "All roles, clinic links and relationships in that organization", "Other organizations, the sign-in account, RehletShifaa coverage"),
        ("End RehletShifaa coverage", "System Administrator", "That staff member's coverage of the organization or clinics", "All provider memberships and clinic links"),
        ("Suspend one clinic", "Authorized provider manager or Provider Operations", "Clinic-specific work at that clinic, for everyone", "The organization and its other clinics"),
        ("Suspend an organization", "Provider Operations Manager", "Work at all of its clinics, for everyone", "Other organizations; history"),
        ("Disable a person's sign-in", "RehletShifaa governance only", "Everything, everywhere", "History, for audit"),
    ], [1.5, 1.6, 2.1, 1.9], 8.6)
    add_label(doc, "Not yet decided:", "a temporary pause of one person's membership in one organization (as opposed to ending it), and whether a provider member can end their own membership themselves. Today the model only defines ending a membership.", RED)
    add_heading(doc, "Internal staff leaving (offboarding)", 2)
    add_numbered(doc, [
        "The person's manager or a System Administrator starts offboarding with a reason. From then on the person receives no new work but keeps enough access to hand over.",
        "If there is a security concern, sign-in is disabled immediately and handover continues without the person.",
        "The platform lists blockers: owned cases, open tasks, travel or finance work, direct reports, and being the last System Administrator.",
        "Authorized leads transfer the work using the normal Transfer and reassignment actions.",
        "When nothing blocks, offboarding completes: access and coverage end, sign-in is disabled, sessions end, and all history is kept. The account is never deleted.",
    ])
    add_heading(doc, "Changing an internal staff member's job", 2)
    add_body(doc, "A job change replaces the old job with the new one in a single step, so the person never has both or neither. It is refused while the person owns work that the new job cannot handle. For example, a Care Coordinator becoming a Finance Officer must first transfer owned cases, while a Care Coordinator becoming a Care Coordination Manager keeps them.")

    # ---- 11 Services
    add_heading(doc, "11 Services, prices, schedules and cases", 1)
    add_heading(doc, "Services and prices", 2)
    add_body(doc, "The organization owns the default service list. A clinic must explicitly offer a service before it can be priced, scheduled or booked there. The platform then uses the most specific active price.")
    add_picture(doc, fig["pricing"], "Five-level price order from clinician-at-clinic to catalogue fallback.", 6.8)
    add_caption(doc, "Figure 9  Price order, from most specific to fallback")
    add_table(doc, ["Prices that exist for a consultation with Dr Sara", "Price charged at Dubai", "Price charged at Sharjah"], [
        ("Organization default 400; Sharjah clinic 450; Dr Sara organization-wide 500; Dr Sara at Dubai 550", "550 (Dr Sara at Dubai)", "450 (Sharjah clinic price beats Dr Sara's organization-wide price)"),
    ], [3.6, 1.6, 1.9], 8.9)
    add_body(doc, "A clinician-specific price must be approved by someone other than its author.")
    add_heading(doc, "Schedules and virtual care", 2)
    add_bullets(doc, [
        "Every availability slot and exception names a clinic. Clinic closures override clinician slots at that clinic.",
        "Virtual care uses the organization's virtual-care clinic. An empty clinic never means virtual.",
        "A clinician cannot be booked in person at two clinics at the same time, or in person and virtually at the same time.",
        "Two virtual slots at the same time are refused unless that virtual clinic allows it (adopted default: not allowed).",
        "A clinician is never copied to represent another clinic, price or schedule.",
    ])
    add_heading(doc, "Cases and routing", 2)
    add_table(doc, ["Moment in the case", "What must be known"], [
        ("A provider Consultant is chosen", "The provider organization is recorded on the case."),
        ("Intake and review", "The clinic may still be open."),
        ("Before an in-person appointment, travel finalization or a location-specific service", "The clinic (adopted default)."),
        ("Care is virtual", "Virtual is recorded explicitly, with the organization's virtual-care clinic."),
        ("A clinician is assigned", "The clinician is active and eligible in the organization and linked to the case's clinic."),
        ("Routing picks a destination", "Routing is by organization at launch; the clinic filters who is eligible (adopted default)."),
    ], [3.0, 4.1], 9.1)

    # ---- 12 Tester handbook
    add_heading(doc, "12 Tester handbook", 1)
    add_callout(doc, "Run these tests on a first-production-release build only", [
        "Every test here is an acceptance test for the first production release. On today's development system many will fail because the feature is not built (see Section 2).",
    ], AMBER_FILL, GOLD)
    add_body(doc, "A role test is incomplete if it checks only success. Every important permission needs a positive test in the right context and negative tests for the wrong organization, clinic, role, status, sign-in strength and qualification.")
    add_picture(doc, fig["test_map"], "Every access decision combines identity, strong sign-in, membership, clinic, eligibility and isolation.", 6.4)
    add_caption(doc, "Figure 10  Vary each condition independently")
    add_heading(doc, "Prerequisites", 2)
    add_bullets(doc, [
        "A build that passes all first-production-release gates, with the sign-in system configured for strong sign-in.",
        "Authenticator secrets for QA accounts kept only in the local environment file, never in source control.",
        "Two separate browser profiles or private windows, so that two people can act at the same time.",
        "Read access to the access audit (a Compliance Auditor account) to collect evidence.",
    ])
    add_heading(doc, "Test world", 2)
    add_table(doc, ["Organization", "Clinics", "Notes"], [
        ("Al Noor Healthcare Group", "Dubai Clinic (default), Sharjah Clinic, Abu Dhabi Hospital, Al Noor Online (virtual)", "Main organization for scope tests"),
        ("Gulf Heart Center", "Gulf Heart Main, Gulf Heart West", "Second organization for isolation tests"),
        ("Crescent Medical", "Crescent Rehab", "Third organization for supervision tests"),
    ], [1.9, 3.3, 1.9], 9.0)
    add_heading(doc, "Minimum test personas", 2)
    add_table(doc, ["Persona", "Setup", "Main purpose"], [
        ("Platform Account Owner", "Active ownership; strong sign-in", "Approvals; no business access"),
        ("System Administrators A and B", "Two active administrators", "Maker and checker; last-administrator rules"),
        ("Care Coordinator", "Covers Al Noor; owns one case", "Case access; job change blocking"),
        ("Provider Operations Specialist", "Covers Al Noor only", "Setup limits; isolation from Gulf Heart"),
        ("Credentialing Specialist (Omar)", "Covers Al Noor Dubai Clinic only", "RehletShifaa coverage tests"),
        ("Organization Owner (Al Noor)", "Organization-wide", "Reach to all clinics; no platform access"),
        ("Organization Owner (Gulf Heart)", "Organization-wide", "Cross-organization isolation"),
        ("Practice Manager, organization-wide", "Al Noor, organization scope", "New-clinic inheritance"),
        ("Practice Manager, clinic-limited", "Al Noor, Dubai only", "Sibling and organization-default denial"),
        ("Dr Sara Haddad", "As in the worked example (Section 13)", "Multi-organization, multi-role"),
        ("Supervising Consultant (Dr Lina Mansour)", "Crescent, Consultant at Crescent Rehab", "Supervision rules"),
        ("Consultant Assistant", "Al Noor, Dubai, assigned to one clinician", "Assistant limits and conflicts"),
        ("Patient and Patient Representative", "Own case; one represented patient", "Regression of patient access"),
        ("Fresh email, internal staff email", "No account; a RehletShifaa employee's email", "Offer and identity tests"),
    ], [2.1, 2.6, 2.4], 8.5)

    tests = [
        ("Identity and offers", [
            ("ID 01", "Two organizations invite the same new email at the same time", "One account and, after acceptance, one practitioner profile. Two separate pending offers."),
            ("ID 02", "Invite a person who already has an account", "A pending offer only. No membership, role or clinic link before acceptance."),
            ("ID 03", "A provider invites a RehletShifaa employee's email", "No offer becomes usable; the case goes to governance review and is refused by default."),
            ("ID 04", "Invite a patient-only person", "Normal offer. On acceptance a practitioner profile is created if needed; Patient access remains."),
            ("ID 05", "Accept the same offer twice", "Nothing is duplicated."),
            ("ID 06", "Decline one offer; let another expire", "Only those offers end. Other offers and memberships remain."),
            ("ID 07", "Send offers to a new email, an existing account and a staff email", "The sender sees exactly the same \"Invitation sent\" response each time."),
            ("ID 08", "Offer Consultant to a person declared as Associate Doctor", "Offer is sent without the Consultant role; sender sees \"some roles need review\"."),
        ]),
        ("Organization and clinic scope", [
            ("SC 01", "Al Noor Owner acts at a clinic created after the grant", "Allowed."),
            ("SC 02", "Organization-wide Practice Manager acts at that new clinic", "Allowed."),
            ("SC 03", "Dubai-only Practice Manager acts at Sharjah", "Denied."),
            ("SC 04", "Dubai-only Practice Manager edits organization defaults", "Denied."),
            ("SC 05", "Suspend Sharjah", "Provider and RehletShifaa clinic-specific actions at Sharjah are denied; Dubai continues."),
            ("SC 06", "Suspend Al Noor", "All Al Noor clinics denied for everyone; Gulf Heart unaffected."),
            ("SC 07", "End a clinician's Sharjah link", "Sharjah removed from their roles; Dubai and other organizations continue."),
            ("SC 08", "Send a clinic identifier that does not match the stored record", "The stored clinic is used; the mismatch is denied."),
            ("SC 09", "Sharjah-only user acts at the default clinic (Dubai)", "Denied. The default clinic grants nothing."),
            ("SC 10", "Activate an organization with no active clinic", "Refused: no active clinic."),
            ("SC 11", "Close the default clinic without naming a replacement", "Refused, or allowed with acknowledgement and a \"default clinic required\" warning."),
        ]),
        ("RehletShifaa coverage", [
            ("CC 01", "Al Noor Owner looks for or tries to remove Omar's coverage", "Not visible in people lists; removal impossible."),
            ("CC 02", "Omar reviews a Dubai clinician's credentials", "Allowed without any provider membership or clinic link."),
            ("CC 03", "Omar opens a Sharjah or Gulf Heart credential file", "Denied."),
            ("CC 04", "End Dr Sara's Dubai link", "Omar's coverage of Dubai is unchanged."),
            ("CC 05", "System Administrator ends Omar's coverage", "No provider membership or clinic link changes."),
        ]),
        ("Roles and eligibility", [
            ("RL 01", "Give Consultant to a person declared as Associate Doctor", "Refused: level insufficient."),
            ("RL 02", "Give Consultant before a Consultant credential decision in this organization", "Refused, even if another organization has verified the person."),
            ("RL 03", "Give Associate Doctor with no supervisor", "Refused: supervision required."),
            ("RL 04", "Consultant at one clinic, Associate Doctor at another", "Allowed for a Consultant-qualified person with supervision for the Associate work."),
            ("RL 05", "Consultant Assistant and Consultant at the same clinic", "Refused."),
            ("RL 06", "Remove one role from a person with several", "Only that role ends."),
            ("RL 07", "Give any clinical role to a person with no declared level", "Refused."),
            ("RL 08", "New Consultant role for a level carried over from the old system", "Refused until the professional confirms the level."),
            ("RL 09", "Organization raises the level it is onboarding a person for", "Credential level stays unchanged until a new independent decision."),
            ("RL 10", "Dr Sara, as Al Noor Owner, gives herself a Consultant role at Abu Dhabi", "Refused: nobody can grant themselves a role."),
        ]),
        ("Separation of duties and administration", [
            ("AD 01", "System Administrator opens a patient document", "Denied."),
            ("AD 02", "Compliance Auditor opens a patient case", "Denied."),
            ("AD 03", "System Administrator grants themselves a role", "Refused."),
            ("AD 04", "The only administrator resigns", "Blocked until a replacement is active."),
            ("AD 05", "Provider Operations person decides credentials of a clinician they invited", "Refused: independent review required."),
            ("AD 06", "Give Care Journey Approver to a Care Journey Manager", "Refused."),
            ("AD 07", "Give System Administrator to a Compliance Auditor", "Refused."),
            ("AD 08", "Practice Manager approves a price they wrote", "Refused."),
            ("AD 09", "Owner submits the same transfer twice", "One pending nomination only."),
            ("AD 10", "A System Administrator tries to transfer ownership", "Refused."),
            ("AD 11", "Care Coordinator owning cases changed to Finance Officer", "Refused; nothing changes. Succeeds after the cases are transferred."),
            ("AD 12", "Care Coordinator owning cases promoted to Care Coordination Manager", "Allowed; cases stay."),
        ]),
        ("Sign-in and lifecycle", [
            ("AU 01", "Owner, administrator or staff member uses a password only", "Strong sign-in required before any business action."),
            ("AU 02", "Accept a manager-role offer without strong sign-in", "Asked to use strong sign-in first."),
            ("AU 03", "A Consultant with no manager role uses a password only", "Strong sign-in required (at launch)."),
            ("AU 04", "A patient uses a password only", "Patient use is unaffected."),
            ("AU 05", "Invited staff member signs in with password only", "Not activated; no access."),
            ("AU 06", "Ten first requests arrive at once for an invited staff member", "Exactly one activation recorded."),
            ("AU 07", "Remove one manager role while another qualifying role remains", "Strong sign-in still required."),
            ("AU 08", "Administrator action with a sign-in older than five minutes", "Asked to sign in again."),
            ("AU 09", "End one provider membership", "Sign-in stays enabled; other memberships continue."),
        ]),
        ("Offboarding", [
            ("OB 01", "Complete offboarding for a coordinator who still owns cases", "Blocked until the cases are transferred."),
            ("OB 02", "Assign a new case to a person being offboarded", "Refused."),
            ("OB 03", "Use the person's old session after offboarding completes", "No access. History is kept."),
            ("OB 04", "Offboard the last System Administrator", "Blocked."),
        ]),
        ("Context and isolation", [
            ("CX 01", "Dr Sara signs in", "Exactly three organizations in the selector; one account; no duplicates."),
            ("CX 02", "Gulf Heart Owner opens an Al Noor page by link", "Denied."),
            ("CX 03", "Gulf Heart manager views Dr Sara's clinician page", "Only her Gulf Heart work is shown."),
            ("CX 04", "Search, notifications and audit for a Gulf Heart manager", "No Al Noor or Crescent data appears."),
            ("CX 05", "\"Can they...?\" check for Dr Sara at Abu Dhabi", "Explains \"No authority at Abu Dhabi Hospital\" in plain words."),
        ]),
        ("Services, prices, schedules and cases", [
            ("OP 01", "Price a consultation where every price level exists", "The clinician-at-clinic price is used."),
            ("OP 02", "Remove the clinician-at-clinic price", "The clinic price is used."),
            ("OP 03", "Price or schedule a service the clinic does not offer", "Denied."),
            ("OP 04", "Overlapping in-person slots at two clinics", "The second is refused, also across time zones."),
            ("OP 05", "Overlapping in-person and virtual slots", "The second is refused."),
            ("OP 06", "In-person appointment step without a clinic", "Refused until an eligible active clinic is chosen."),
            ("OP 07", "Assign a clinician not linked to the case's clinic", "Refused."),
            ("OP 08", "Case with no delivery mode set", "Not treated as virtual."),
            ("OP 09", "Open a case of another organization or clinic", "Denied."),
        ]),
    ]
    for heading, rows in tests:
        add_heading(doc, heading, 3)
        add_table(doc, ["Test", "Action", "Expected result"], rows, [0.75, 3.05, 3.3], 8.5, BLUE)
    add_heading(doc, "Evidence to keep for every test", 2)
    add_bullets(doc, [
        "Who was signed in and whether strong sign-in was used.",
        "The organization, clinic, role and scope involved.",
        "Membership, clinic link, role and eligibility status at the time.",
        "The response and the plain-language reason for allow or deny.",
        "The audit entry: who, for whom, organization, clinic, before and after, reason and reference numbers.",
        "Proof that unrelated organizations, clinics, roles and accounts did not change.",
    ])

    # ---- 13 Worked example
    add_heading(doc, "13 Worked example", 1)
    add_heading(doc, "Meet Dr Sara", 2)
    add_body(doc, "Dr Sara Haddad has one sign-in and one practitioner profile. She has declared herself a Consultant. Each organization still gets its own independent credential decision and grants its own local roles.")
    add_picture(doc, fig["example"], "Dr Sara is Owner and Consultant at Al Noor, Practice Manager at Gulf Heart Main and Associate Doctor at Crescent Rehab.", 6.8)
    add_caption(doc, "Figure 11  One person, different roles in different places")
    add_table(doc, ["Organization", "Clinics", "Roles", "Why access works"], [
        ("Al Noor Healthcare Group", "Dubai and Sharjah", "Organization Owner (whole organization); Consultant at Dubai and Sharjah", "Active membership and clinic links; Consultant declaration; Consultant credential decision for Al Noor"),
        ("Gulf Heart Center", "Gulf Heart Main", "Practice Manager at Gulf Heart Main", "Active membership and clinic link; clinic-limited manager role"),
        ("Crescent Medical", "Crescent Rehab", "Associate Doctor", "Active membership and clinic link; credential decision for Crescent at Associate level or higher; supervised by Dr Lina Mansour"),
    ], [1.45, 1.2, 2.2, 2.25], 8.5)
    add_heading(doc, "What Dr Sara may and may not do", 2)
    add_table(doc, ["Question", "Answer", "Reason"], [
        ("Can she manage Al Noor's Sharjah Clinic as owner?", "Yes", "The Owner role covers every Al Noor clinic."),
        ("Can she manage Abu Dhabi Hospital as owner?", "Yes", "Owner reach includes every clinic, even ones she is not linked to."),
        ("Can she consult at Abu Dhabi Hospital?", "Not yet", "Her Consultant role names only Dubai and Sharjah. Another authorized Al Noor manager must add a clinic link and extend the role; she cannot do it for herself."),
        ("Can she manage Gulf Heart West?", "No", "Her manager role names only Gulf Heart Main."),
        ("Can she work as a Consultant at Crescent Rehab?", "No", "Her Crescent role is Associate Doctor, and Crescent has made no Consultant-level decision for her."),
        ("Can she approve her own credentials because she is an owner?", "No", "Self-review is always refused."),
        ("Can she make someone a RehletShifaa System Administrator?", "No", "Provider authority never reaches platform governance."),
        ("Can she sign in with a password only when working at Crescent?", "No", "Strong sign-in follows the person, not the context."),
    ], [2.5, 0.9, 3.7], 8.7)
    add_heading(doc, "What changes when something ends", 2)
    add_numbered(doc, [
        "Gulf Heart removes her Practice Manager role: only that role ends. Al Noor and Crescent are unchanged.",
        "Al Noor ends her Sharjah clinic link: she can no longer consult at Sharjah, but as Organization Owner she can still manage Sharjah. Dubai consulting continues.",
        "Crescent's supervision relationship ends: her Associate Doctor work at Crescent pauses until supervision is restored. Other organizations are unchanged.",
        "Gulf Heart suspends Gulf Heart Main: work at Main stops for everyone there. Al Noor and Crescent are unchanged.",
        "RehletShifaa disables her sign-in after a confirmed account compromise: everything stops until the global action is resolved.",
    ])
    add_heading(doc, "Example test script", 2)
    add_table(doc, ["Step", "Action", "Expected result"], [
        ("1", "Sign in as Dr Sara with strong sign-in", "Three organizations in the selector; one account."),
        ("2", "Choose Al Noor, then Dubai Clinic", "Owner and Consultant actions allowed, each by its own rules."),
        ("3", "Switch to Sharjah Clinic", "Consultant work allowed through the Sharjah link."),
        ("4", "Switch to Abu Dhabi Hospital", "Owner management allowed; consulting denied."),
        ("5", "Switch to Gulf Heart Main", "Practice Manager actions allowed; no clinical role there."),
        ("6", "Open a Gulf Heart West page by link", "Denied."),
        ("7", "Switch to Crescent Rehab", "Associate Doctor work allowed while supervision is active."),
        ("8", "Remove the Gulf Heart manager role; repeat steps 2 and 7", "Only Gulf Heart access has ended."),
    ], [0.6, 3.1, 3.4], 8.7)

    # ---- 14 Quick reference
    add_heading(doc, "14 Quick reference, glossary and decision status", 1)
    add_heading(doc, "Role catalogue at a glance", 2)
    add_table(doc, ["Group", "Roles"], [
        ("Platform governance", "Platform Account Owner; System Administrator"),
        ("RehletShifaa internal staff", "Care Coordination Manager; Care Coordinator; Travel and Logistics Manager; Travel and Logistics Specialist; Commercial and Finance Manager; Finance Officer; Provider Operations Manager; Provider Operations Specialist; Credentialing Specialist; Patient Identity Reviewer; Care Journey Manager; Care Journey Approver; Compliance Auditor"),
        ("Provider organization", "Organization Owner; Practice Manager; Consultant; Associate Doctor; Consultant Assistant"),
        ("Patients", "Patient; Patient Representative"),
    ], [1.8, 5.3], 9.0)
    add_heading(doc, "Glossary", 2)
    add_table(doc, ["Term", "Meaning"], [
        ("Platform Account Owner", "Owns the RehletShifaa platform account. Not a provider role."),
        ("Organization Owner", "Owns one provider organization. Has no platform governance authority."),
        ("Provider organization", "The legal or contracting provider party, and the boundary between providers."),
        ("Clinic or facility", "A clinic, branch, hospital, center or virtual-care site inside one organization. Shown as \"Clinics & branches\"."),
        ("Default clinic", "The one clinic per organization that pre-fills forms. Grants nothing."),
        ("Membership", "A provider person's relationship with one organization, with its own start and end."),
        ("Clinic link (affiliation)", "Where a provider member works. Grants nothing by itself."),
        ("RehletShifaa coverage", "A RehletShifaa employee's assignment to serve an organization or clinics without becoming provider staff."),
        ("Role", "What a person may do."),
        ("Scope", "Where the role applies: the organization, selected clinics, managed clinicians, self, or the platform."),
        ("Offer", "An invitation to join an organization that the person must accept."),
        ("Declared level", "The professional's own statement: Consultant or Associate Doctor."),
        ("Credential decision", "The level verified for one organization by an independent RehletShifaa Credentialing Specialist."),
        ("Strong sign-in", "Sign-in with a second factor such as an authenticator app. Some actions also need it within the last five minutes."),
        ("First production release (P1)", "The single release that brings all of this model live together."),
    ], [2.0, 5.1], 9.0)
    add_heading(doc, "Before allowing any protected action", 2)
    add_numbered(doc, [
        "Is the person's sign-in enabled?",
        "For RehletShifaa staff: is the staff record active (or in offboarding handover)?",
        "Is the provider membership or RehletShifaa coverage active in this organization?",
        "Is the organization active?",
        "Is the clinic active, when the record belongs to a clinic?",
        "For a provider role limited to clinics: is the person's link to this clinic active?",
        "Is the role active, and does its scope cover this record?",
        "Was strong sign-in used, and recently enough, when required?",
        "For clinical work: are the declared level, credential decision and supervision valid?",
        "Are self-approval and separation-of-duty rules respected?",
    ])
    add_heading(doc, "Decision status", 2)
    add_body(doc, "The architecture lists fifteen business choices. Two are explicitly confirmed, one is still pending, and the rest stand as adopted defaults.")
    add_table(doc, ["#", "Business choice", "Status", "What this guide assumes"], [
        ("1", "Final English and Arabic names, including clinic terms", "Pending native-language review (wording only)", "Current English working names"),
        ("6", "Strong sign-in for clinicians", "Confirmed", "Required at launch"),
        ("9", "Several clinics and several organizations at launch", "Confirmed", "Yes, in the first production release"),
        ("2", "\"Travel and Logistics\" as the travel team's name", "Adopted default", "As named"),
        ("3", "Commercial and Finance as one function", "Adopted default", "Manager plus Officer"),
        ("4", "Recovery waiting period and request expiry", "Adopted default", "72 hours each"),
        ("5", "Platform owner may also be an administrator", "Adopted default", "Only via another administrator's request and approval"),
        ("7", "Extra-responsibility set and conflicts", "Adopted default", "As in Section 8"),
        ("8", "Second-factor types", "Adopted default", "Authenticator app required; security keys optional"),
        ("10", "Consultant-qualified person as Associate elsewhere", "Adopted default", "Allowed with supervision, not at the same clinic"),
        ("11", "RehletShifaa employee as provider member", "Adopted default", "Refused"),
        ("12", "Offer expiry", "Adopted default", "14 days"),
        ("13", "Concurrent virtual slots", "Adopted default", "Not allowed unless the virtual clinic opts in"),
        ("14", "When a case needs a clinic", "Adopted default", "Before in-person care, travel finalization or location-specific service"),
        ("15", "Clinic-aware routing at launch", "Adopted default", "Routing by organization; clinic as eligibility filter"),
    ], [0.35, 2.55, 1.75, 2.45], 8.4)
    add_heading(doc, "Not yet decided (not covered by the architecture)", 2)
    add_table(doc, ["Question", "Why it matters"], [
        ("Who may grant and remove the Organization Owner role, must every organization keep at least one active owner, and how does organization ownership pass to someone else?", "Without a rule, an organization could lose its last owner, or ownership could change without a clear approver."),
        ("Can one person's membership be paused temporarily in one organization, and can a provider member end their own membership?", "The model defines only ending a membership by an authorized manager."),
    ], [4.2, 2.9], 9.0)
    add_heading(doc, "Final rule", 2)
    add_body(doc, "A person can have several legitimate relationships without gaining broad or accidental access. RehletShifaa keeps one identity per person, records every organization and clinic relationship separately, checks each role in its exact context, and changes only the relationship that an authorized action names.")
    add_body(doc, "Source: Customer Governance Target Architecture (Reviewed), revision 3 (final), decisions D1 to D35 and business choices 1 to 15. Edition 2 corrects Edition 1; the review of Edition 1 is kept alongside this document.", italic=True)

    props = doc.core_properties
    props.title = "RehletShifaa Users, Roles and Access Guide (Edition 2)"
    props.subject = "Platform and provider roles, scopes, lifecycle and tester guidance"
    props.author = "RehletShifaa"
    props.comments = "Target operating model for the first production release; not yet implemented"
    doc.save(DOCX)
    return DOCX


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    fig = {
        "populations": diagram_populations(),
        "governance": diagram_governance(),
        "ownership": diagram_ownership_transfer(),
        "provider_tree": diagram_provider_tree(),
        "multi_clinic": diagram_multi_clinic(),
        "two_paths": diagram_two_paths(),
        "context": diagram_context_switch(),
        "offer": diagram_offer_flow(),
        "lifecycle": diagram_lifecycle(),
        "pricing": diagram_pricing(),
        "example": diagram_example(),
        "test_map": diagram_test_map(),
    }
    print(build_doc(fig))


if __name__ == "__main__":
    main()
